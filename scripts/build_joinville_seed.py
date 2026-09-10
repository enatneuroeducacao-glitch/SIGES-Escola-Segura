from pathlib import Path
import re
import math
import openpyxl

INPUT = Path('HSI_DOTH_P_ESCOLAR_JOINVILLE_3_2_ALUNO_GUIA.xlsx')
OUT = Path('supabase/migrations/0012_joinville_162_seed_generated.sql')
SHEET = 'MATRIZ 3.2'

CORE = ['name','unit_type','neighborhood','address','rank','d','o','t','h','p','hsi','infrastructure_score','exposure_score','technical_evidence_score','public_evidence_score','confidence_score','ipe','priority','student_guide','behavior_flag','traffic_flag','physical_infrastructure_flag','transport_crossing_flag','accessibility_flag','speed_kmh','vdm','crashes_3y','crossings','sidewalk','signage','lighting','pickup_dropoff','cyclists','primary_source','data_status','recommended_action','observations']
TERR = ['corridor_normalized','detrans_in_corridor','detrans_studies_count','detrans_proximity','detrans_point','accidents_corridor_2023','accidents_corridor_2024','accident_rank_2024','accident_evidence_score','detrans_evidence_score','territorial_data_available','territorial_confidence','ipe_territorial','territorial_priority','territorial_justification','territorial_primary_source']
COLS = CORE + TERR

def key(s):
    s = s.lower().replace('ã','a').replace('ç','c').replace('ê','e').replace('é','e').replace('í','i').replace('õ','o').replace('á','a').replace('ú','u')
    return re.sub(r'[^a-z0-9]+','_',s).strip('_')

def sql(v):
    if v is None or v == '': return 'NULL'
    if isinstance(v, (int,float)) and not isinstance(v,bool):
        if isinstance(v,float) and (math.isnan(v) or math.isinf(v)): return 'NULL'
        return str(v)
    return "'" + str(v).replace("'", "''") + "'"

wb = openpyxl.load_workbook(INPUT, data_only=True)
ws = wb[SHEET]
headers = [c.value for c in ws[1]]
rows = list(ws.iter_rows(min_row=2, values_only=True))
if len(rows) != 162:
    raise SystemExit(f'Esperadas 162 unidades; encontradas {len(rows)}.')
source = {key(h): h for h in headers}
values = []
for row in rows:
    record = {key(h): v for h,v in zip(headers,row)}
    values.append('(' + ','.join(sql(record.get(c)) for c in COLS) + ')')

updates = ','.join(f'{c}=excluded.{c}' for c in COLS[1:])
text = f"insert into public.units ({','.join(COLS)}) values\n{',\n'.join(values)}\non conflict (name) do update set {updates};\n"
OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(text, encoding='utf-8')
print(f'Gerado: {OUT} ({len(text):,} caracteres)')
